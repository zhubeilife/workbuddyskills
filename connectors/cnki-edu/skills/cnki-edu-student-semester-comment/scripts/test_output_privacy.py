# -*- coding: utf-8 -*-
"""交付物隐私单测（2026-10-12 用户要求：交付文件不要暴露接口/生成侧的元数据字段值）。

交付物只有一份 excel，它最容易被忽略的外露点就是**文件属性**：
  · `docProps/core.xml` 的 `<dc:creator>` —— openpyxl 默认写死 "openpyxl"，
    老师在「文件 → 信息 → 属性」里会看到「作者：openpyxl」
  · `docProps/app.xml` 的 `<Application>` —— openpyxl 写死成
    "Microsoft Excel Compatible / Openpyxl 3.1.5"，把生成库与**精确版本号**带出去

本文件把这些钉死，改坏了立刻红：

  A 单元 · scrub 的零副作用与幂等
    A1 交付 xlsx 的 core.xml 里没有 creator=openpyxl
    A2 app.xml 里没有库名与版本号
    A3 整包任何条目都不含生成侧串
    A4 scrub **只改 docProps 两个条目**，其余条目字节级不变
    A5 scrub 幂等（二次调用命中为空）
    A6 表内容/合并区/列宽/行高/冻结窗格零改动
    A7 检测器本身有效（未清理的裸 openpyxl 文件必须被检出——防"永远通过"的空闸门）
  B 单元 · 真实链路出表（造材料 → 对齐 → 出表）
    B1 真实产物干净（无生成侧串）
    B2 控制台打出「文档属性：已清空」
    B3 交付表仍是一张表、两列（姓名 + 评语），没有多出内部字段列
    B4 正文不含内部字段名（roster.json / comments.json / 学号 / 序号 / 证据来源 /
       分层 / 核心印象标签 / 字数 / json …）
  C 单元 · 技能文件自身
    C1 py/md/json 里不写死个人绝对路径
    C2 不带外链与凭据
    C3 已发布示例 references/期末评语表.xlsx 同样干净
    C4 文档写明这条口径

跑法（用装了 openpyxl 的解释器）：
  python test_output_privacy.py
"""
import glob
import json
import os
import shutil
import sys
import tempfile
import zipfile

import hashlib

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import export_comments as EC        # noqa: E402

PREP = os.path.join(HERE, "prepare_roster.py")
EXPORT = os.path.join(HERE, "export_comments.py")
MSD = os.path.join(HERE, "make_sample_data.py")
SKILL_DIR = os.path.dirname(HERE)
EXAMPLES = os.path.join(SKILL_DIR, "references")
PY = sys.executable

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print("%s %s%s" % ("✓" if cond else "✗", name,
                       ("  → " + str(detail)) if (detail and not cond) else ""))


def run(cmd):
    import subprocess
    p = subprocess.run([PY] + cmd, capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def zip_hashes(path):
    with zipfile.ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}


def sheet_facts(path, name="评语表"):
    wb = load_workbook(path)
    ws = wb[name]
    return {
        "cells": [[c.value for c in row] for row in ws.iter_rows()],
        "merged": sorted(str(r) for r in ws.merged_cells.ranges),
        "widths": {k: v.width for k, v in ws.column_dimensions.items()},
        "heights": {k: v.height for k, v in ws.row_dimensions.items()},
        "freeze": ws.freeze_panes,
        "sheets": list(wb.sheetnames),
    }


def cell_text(path, name="评语表"):
    wb = load_workbook(path)
    ws = wb[name]
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)


# 正文里不该出现的内部字段名（交付表只有「姓名 + 评语 + 小贴士」）
INTERNAL_KEYS = ("roster.json", "comments.json", "analysis.json", "data_notes",
                 "证据来源", "核心印象标签", "分层", "字数", "序号", "学号",
                 "meta", "json", "openpyxl", "WorkBuddy", "Skill", "skill")


def main():
    tmp = tempfile.mkdtemp(prefix="privacy_")
    try:
        # ---------------- A 单元 · scrub 的零副作用与幂等 ----------------
        print("=" * 68)
        print("A 单元 · scrub 只清 docProps，表内容零改动（幂等）")
        src = os.path.join(tmp, "a_bare.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "评语表"
        ws.append(["姓名", "评语"])
        ws.append(["张三", "课堂上你举手的那个瞬间，老师就知道问题要往深里走了。"])
        ws.append(["李四", "早读的铃声还没停，你手里那本书已经翻开了。"])
        wb.save(src)
        before = zip_hashes(src)
        before_facts = sheet_facts(src)

        # A7 先证明检测器不是空闸门：裸文件必须被检出
        bare_leaks = EC.workbook_prop_leaks(src)
        check("A7 检测器有效：未清理的裸 openpyxl 文件必须被检出（防空闸门）",
              bare_leaks != [], bare_leaks)
        check("A7b 裸文件确实带 creator=openpyxl",
              (load_workbook(src).properties.creator or "") == "openpyxl",
              load_workbook(src).properties.creator)

        hits = EC.scrub_xlsx_props(src)
        check("A1 清理命中 docProps（core.xml + app.xml）",
              sorted(hits) == ["docProps/app.xml", "docProps/core.xml"], hits)
        check("A1b 清理后整包无生成侧串", EC.workbook_prop_leaks(src) == [],
              EC.workbook_prop_leaks(src))
        core = zipfile.ZipFile(src).read("docProps/core.xml").decode("utf-8")
        app = zipfile.ZipFile(src).read("docProps/app.xml").decode("utf-8")
        check("A2 core.xml 的 creator 已清空（无 openpyxl）",
              "openpyxl" not in core.lower(), core)
        check("A2b app.xml 无库名与版本号",
              "openpyxl" not in app.lower()
              and "microsoft excel compatible" not in app.lower(), app)
        after = zip_hashes(src)
        changed = [n for n in before if before[n] != after.get(n)]
        check("A4 scrub 只改 docProps，其余条目字节级不变",
              all(n.startswith("docProps/") for n in changed)
              and sorted(before) == sorted(after), changed)
        check("A5 幂等：二次 scrub 命中为空", EC.scrub_xlsx_props(src) == [])
        after_facts = sheet_facts(src)
        for lb, k in (("单元格内容", "cells"), ("合并区", "merged"),
                      ("列宽", "widths"), ("行高", "heights"),
                      ("冻结窗格", "freeze"), ("工作表", "sheets")):
            check("A6 scrub 后 %s 零改动" % lb, before_facts[k] == after_facts[k])
        check("A6b 清理后文件仍可正常打开",
              load_workbook(src)["评语表"]["A2"].value == "张三")

        # ---------------- B 单元 · 真实链路出表 ----------------
        print("=" * 68)
        print("B 单元 · 真实链路（造材料 → 对齐 → 出表）的产物")
        p = lambda n: os.path.join(tmp, n)
        code, log = run([MSD, "--out", tmp])
        check("B0 示例材料生成成功", code == 0, log[-300:])
        code, log = run([PREP,
                         "--scores", p("示例_成绩表.xlsx"),
                         "--homework", p("示例_作业统计表.xlsx"),
                         "--classroom", p("示例_课堂表现统计表.xlsx"),
                         "--notes", p("示例_教师个别描述.xlsx"),
                         "--past-comments", p("示例_往年评语.xlsx"),
                         "--awards", p("示例_奖状荣誉.xlsx"),
                         "--class-name", "初二(3)班",
                         "--exam-name", "初二上学期期末考试",
                         "--out", p("roster.json")])
        check("B0b 材料对齐成功", code == 0, log[-300:])
        shutil.copy(os.path.join(EXAMPLES, "comments.json"), p("comments.json"))
        code, log = run([EXPORT, "--roster", p("roster.json"),
                         "--comments", p("comments.json"),
                         "--out", p("评语表.xlsx"),
                         "--min-len", "120", "--max-len", "180"])
        check("B0c 出表成功（无硬错误）", code == 0 and "硬错误" not in log, log[-400:])
        out = p("评语表.xlsx")
        check("B1 真实产物整包无生成侧串", EC.workbook_prop_leaks(out) == [],
              EC.workbook_prop_leaks(out))
        check("B2 控制台提示文档属性已清空", "文档属性：已清空" in log, log[-400:])
        wb = load_workbook(out)
        check("B3 仍是一张表叫「评语表」", wb.sheetnames == ["评语表"], wb.sheetnames)
        ws = wb["评语表"]
        check("B3b 表头仍是「姓名 + 评语」两列",
              [ws.cell(row=1, column=1).value, ws.cell(row=1, column=2).value]
              == ["姓名", "评语"])
        check("B3c 没有多出内部字段列（第三列为空）",
              all(ws.cell(row=r, column=3).value in (None, "")
                  for r in range(1, ws.max_row + 1)))
        txt = cell_text(out)
        hit = [k for k in INTERNAL_KEYS if k in txt]
        check("B4 正文不含内部字段名", hit == [], hit)
        check("B4b 保留表下小贴士（缺料时提示补材料）", "小贴士" in txt)
        check("B4c 交付表里有 40 条评语", ws.max_row == 45, ws.max_row)

        # ---------------- C 单元 · 技能文件自身 ----------------
        print("=" * 68)
        print("C 单元 · 技能文件自身不写死个人路径 / 不带凭据")
        files = []
        for pat in ("*.py", "*.md", "*.json"):
            files += glob.glob(os.path.join(HERE, pat))
            files += glob.glob(os.path.join(SKILL_DIR, pat))
            files += glob.glob(os.path.join(SKILL_DIR, "references", pat))
            files += glob.glob(os.path.join(EXAMPLES, pat))
        files = sorted(set(files))
        check("C0 收集到技能文件", len(files) >= 20, len(files))
        abs_path_hits, cred_hits = [], []
        import re
        cred_re = re.compile(r"(https?://\S+)|(token\s*[=:])|(secret\s*[=:])"
                             r"|(appid\s*[=:])|(password\s*[=:])|(sk-[A-Za-z0-9]{10,})",
                             re.I)
        path_re = re.compile(r"([A-Za-z]:[\\/]Users[\\/])|(/Users/[A-Za-z0-9_.-]+/)")
        for f in files:
            try:
                t = open(f, encoding="utf-8").read()
            except Exception:
                continue
            for m in path_re.finditer(t):
                abs_path_hits.append("%s: %s" % (os.path.relpath(f, SKILL_DIR), m.group(0)))
            for m in cred_re.finditer(t):
                cred_hits.append("%s: %s" % (os.path.relpath(f, SKILL_DIR), m.group(0)[:60]))
        check("C1 不写死个人绝对路径（Windows/macOS 用户目录）",
              abs_path_hits == [], abs_path_hits[:5])
        check("C2 不带外链与凭据（token/secret/appid/password/sk-）",
              cred_hits == [], cred_hits[:5])
        ex_xlsx = os.path.join(EXAMPLES, "期末评语表.xlsx")
        check("C3 已发布示例 references/期末评语表.xlsx 干净",
              os.path.exists(ex_xlsx) and EC.workbook_prop_leaks(ex_xlsx) == [],
              EC.workbook_prop_leaks(ex_xlsx) if os.path.exists(ex_xlsx) else "文件缺失")
        docs = {}
        for r in ("SKILL.md", "references/data-spec.md", "references/writing-rules.md"):
            fp = os.path.join(SKILL_DIR, r)
            docs[r] = open(fp, encoding="utf-8").read() if os.path.exists(fp) else ""
        check("C4 SKILL.md 写明「交付物不带生成侧元数据」的口径",
              "文档属性" in docs["SKILL.md"] or "元数据" in docs["SKILL.md"],
              "SKILL.md 未写该口径")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print("=" * 68)
    print("通过 %d / 失败 %d" % (len(PASS), len(FAIL)))
    if FAIL:
        print("失败用例：" + "、".join(FAIL))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
